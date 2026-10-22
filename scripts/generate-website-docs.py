"""Build the bilingual HTML and Word manuals from reviewed Markdown sources.

Run with the Python supplied by Codex workspace dependencies (python-docx).
The .md sources are authoritative; this script never reads .env or a database.
"""
import argparse
import html
import re
import shutil
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
STEMS = [f"ASLUNG_{kind}_{lang}" for kind in ("Website_User_Manual", "API_User_Guide") for lang in ("ZH", "EN")]


def blocks(source):
    lines = source.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            language = line[3:]
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            yield "code", language, "\n".join(code)
        elif line.startswith("|"):
            table = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in row):
                    table.append(row)
                i += 1
            yield "table", table
            continue
        elif match := re.match(r"^(#{1,3}) (.*)", line):
            yield "heading", len(match[1]), match[2]
        elif match := re.match(r"^(- |\d+\. )(.*)", line):
            yield "list", match[1], match[2]
        else:
            yield "paragraph", line
        i += 1


def plain(text):
    return re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1", text).replace("**", "").replace("`", "")


def inline(text):
    value = html.escape(text)
    value = re.sub(r"`([^`]+)`", r"<code>\1</code>", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", value)
    return re.sub(r"\[([^]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', value)


def web_document(stem, source):
    parsed = list(blocks(source))
    title = plain(next(block[2] for block in parsed if block[0] == "heading"))
    zh = stem.endswith("ZH")
    toc = []
    content = []
    chapter = 0
    for block in parsed:
        if block[0] == "heading":
            level, text = block[1:]
            chapter += 1
            anchor = f"section-{chapter}"
            content.append(f'<h{level} id="{anchor}">{inline(text)}</h{level}>')
            if level == 2:
                toc.append(f'<li><a href="#{anchor}">{inline(text)}</a></li>')
        elif block[0] == "paragraph":
            content.append(f"<p>{inline(block[1])}</p>")
        elif block[0] == "list":
            prefix = "• " if block[1] == "- " else block[1]
            content.append(f'<p class="step">{prefix}{inline(block[2])}</p>')
        elif block[0] == "code":
            content.append(f'<pre><code>{html.escape(block[2])}</code></pre>')
        elif block[0] == "table":
            rows = block[1]
            head = "".join(f"<th>{inline(cell)}</th>" for cell in rows[0])
            body = "".join("<tr>" + "".join(f"<td>{inline(cell)}</td>" for cell in row) + "</tr>" for row in rows[1:])
            content.append(f'<div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>')
    label = ("下載 Word", "下載 Markdown", "列印", "內容導覽") if zh else ("Download Word", "Download Markdown", "Print", "Contents")
    sibling = stem[:-2] + ("EN" if zh else "ZH")
    return f'''<!doctype html>
<html lang="{'zh-Hant' if zh else 'en'}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)}</title>
<style>
*{{box-sizing:border-box}}html{{scroll-behavior:smooth}}body{{margin:0;background:#f5f7fa;color:#183044;font-family:system-ui,"Microsoft JhengHei",sans-serif;line-height:1.85}}header{{position:sticky;top:0;background:white;border-bottom:1px solid #d9e2ea;padding:12px 24px;z-index:2;display:flex;gap:18px;flex-wrap:wrap;align-items:center}}a{{color:#166a67}}header a,button{{font-size:14px}}button{{border:1px solid #bfcbd5;border-radius:6px;background:white;padding:5px 12px;cursor:pointer}}.layout{{display:grid;grid-template-columns:250px minmax(0,900px);gap:28px;max-width:1240px;margin:28px auto;padding:0 20px}}nav{{font-size:13px;position:sticky;top:90px;align-self:start;max-height:80vh;overflow:auto}}nav ul{{padding-left:18px}}nav li{{margin:8px 0}}main{{background:white;padding:36px 44px;border:1px solid #e0e7ee;border-radius:10px;min-width:0}}h1,h2,h3{{color:#000;line-height:1.45}}h1{{font-size:30px;margin-top:0}}h2{{font-size:23px;margin-top:36px;padding-top:10px;scroll-margin-top:100px}}h3{{font-size:18px}}p{{overflow-wrap:anywhere}}.step{{margin-left:14px}}.table-wrap{{overflow:auto;margin:20px 0}}table{{border-collapse:collapse;width:100%;font-size:14px}}th,td{{border:1px solid #d9d9d9;padding:10px 12px;text-align:left;vertical-align:middle}}th{{background:#eaf0f5;color:black}}tbody tr:nth-child(even){{background:#fafbfc}}pre{{overflow:auto;padding:18px;border:1px solid #d9e2ea;border-radius:6px;font-size:12px;line-height:1.65;background:#f7f9fb}}code{{font-family:Consolas,monospace}}footer{{margin-top:36px;font-size:12px;color:#63788a}}@media(max-width:900px){{.layout{{display:block}}nav{{position:static;max-height:none;margin-bottom:20px}}main{{padding:24px 18px}}h1{{font-size:25px}}}}@media print{{header,nav{{display:none}}.layout{{display:block;margin:0;padding:0}}main{{border:0;padding:0}}body{{background:white;color:black}}pre{{white-space:pre-wrap}}h2,h3{{break-after:avoid}}tr{{break-inside:avoid}}}}
</style></head><body><header><strong>ASLUNG</strong><a href="{stem}.docx" download>{label[0]}</a><a href="{stem}.md" download>{label[1]}</a><a href="{sibling}.html">{'English' if zh else '繁體中文'}</a><button onclick="window.print()">{label[2]}</button></header><div class="layout"><nav aria-label="{label[3]}"><strong>{label[3]}</strong><ul>{''.join(toc)}</ul></nav><main>{''.join(content)}<footer>ASLUNG · 2026-10-22 · {'請依部署環境及實際授權操作。' if zh else 'Use the deployed configuration and actual permissions.'}</footer></main></div></body></html>
'''


def word_document(stem, source, target):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    font = "Microsoft JhengHei" if stem.endswith("ZH") else "Calibri"
    for name, size in [("Normal", 11), ("Title", 24), ("Heading 1", 17), ("Heading 2", 13)]:
        style = document.styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft JhengHei")
        style.paragraph_format.line_spacing = 1.10
        style.paragraph_format.space_after = Pt(6)
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
    normal = document.styles["Normal"]
    normal.paragraph_format.widow_control = True
    for name in ("Title", "Heading 1", "Heading 2"):
        properties = document.styles[name].element.get_or_add_pPr()
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
    header = section.header.paragraphs[0]
    header.text = "ASLUNG"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for block in blocks(source):
        kind = block[0]
        if kind == "heading":
            document.add_paragraph(plain(block[2]), "Title" if block[1] == 1 else f"Heading {block[1]-1}")
        elif kind == "paragraph":
            document.add_paragraph(plain(block[1]))
        elif kind == "list":
            prefix = "• " if block[1] == "- " else block[1]
            paragraph = document.add_paragraph(prefix + plain(block[2]))
            paragraph.paragraph_format.left_indent = Inches(.12)
        elif kind == "code":
            code_lines = block[2].splitlines()
            for index, line in enumerate(code_lines):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                # Keep logical code paragraphs together, but allow long examples
                # to break at blank lines rather than leaving half a page empty.
                paragraph.paragraph_format.keep_with_next = (
                    bool(line) and index < len(code_lines) - 1 and bool(code_lines[index + 1])
                )
                run = paragraph.add_run(line or " ")
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            document.add_paragraph().paragraph_format.space_after = Pt(3)
        elif kind == "table":
            rows = block[1]
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = [2, 5] if len(rows[0]) == 2 else [1.7, 2.65, 2.65]
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            borders = OxmlElement("w:tblBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                edge = OxmlElement("w:" + side)
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    edge.set(qn("w:" + key), value)
                borders.append(edge)
            table._tbl.tblPr.append(borders)
            for number, values in enumerate(rows):
                cells = table.add_row().cells
                # Keep a row together so a trailing word is not stranded on the next page.
                table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for cell, width, value in zip(cells, widths, values):
                    cell.width = Inches(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = plain(value)
                    properties = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "left", "bottom", "right"):
                        margin = OxmlElement("w:" + side)
                        margin.set(qn("w:w"), "90")
                        margin.set(qn("w:type"), "dxa")
                        margins.append(margin)
                    properties.append(margins)
                    if number == 0 or number % 2 == 0:
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "EAF0F5" if number == 0 else "F7F9FB")
                        properties.append(shade)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(3)
                        paragraph.paragraph_format.space_before = Pt(3)
                        paragraph.paragraph_format.keep_with_next = number == 0
                        for run in paragraph.runs:
                            run.font.size = Pt(10)
                            run.bold = number == 0
                if number == 0:
                    repeat = OxmlElement("w:tblHeader")
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            document.add_paragraph().paragraph_format.space_after = Pt(3)
    document.core_properties.title = plain(next(item[2] for item in blocks(source) if item[0] == "heading"))
    document.core_properties.subject = "ASLUNG website documentation revision 2026-10-22"
    document.core_properties.author = "ASLUNG"
    document.save(target)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--publish-only", action="store_true", help="Republish native-Word-updated DOCX without rebuilding")
    args = parser.parse_args()
    output = ROOT / "public" / "docs"
    output.mkdir(parents=True, exist_ok=True)
    for stem in STEMS:
        source_path = ROOT / "docs" / (stem + ".md")
        source = source_path.read_text(encoding="utf-8")
        web = web_document(stem, source)
        target_docx = ROOT / "docs" / (stem + ".docx")
        if args.check:
            assert (output / (stem + ".md")).read_text(encoding="utf-8") == source, stem
            assert (output / (stem + ".html")).read_text(encoding="utf-8") == web, stem
            assert (output / (stem + ".docx")).read_bytes() == target_docx.read_bytes(), stem
        else:
            if not args.publish_only:
                word_document(stem, source, target_docx)
            (output / (stem + ".md")).write_text(source, encoding="utf-8")
            (output / (stem + ".html")).write_text(web, encoding="utf-8")
            shutil.copyfile(target_docx, output / target_docx.name)
        print(("Checked " if args.check else "Published ") + stem)


if __name__ == "__main__":
    main()
